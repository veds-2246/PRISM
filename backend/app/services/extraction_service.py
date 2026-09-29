from dataclasses import dataclass
from functools import lru_cache
import io
import re
import subprocess
import tempfile


@dataclass
class ExtractedDocument:
    text: str
    ocr_used: bool = False
    ocr_confidence: float | None = None
    pages: list[str] | None = None


def extract_text(content: bytes, extension: str) -> ExtractedDocument:
    extension = extension.lower()
    if extension in {".txt", ".csv"}:
        text = content.decode("utf-8", errors="replace")
        if not text.strip():
            raise ValueError("The text file is empty")
        return ExtractedDocument(text)
    if extension == ".pdf":
        if not content:
            raise ValueError("The PDF file is empty")
        try:
            from pypdf import PdfReader
            try:
                pages = [page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages]
                text = "\n".join(pages)
            except Exception as exc:
                raise ValueError("The PDF could not be read; it may be corrupted") from exc
            return _extract_pdf_pages_with_ocr(content, pages)
        except ImportError as exc:
            raise RuntimeError("PDF support is not installed") from exc
    if extension == ".docx":
        try:
            from docx import Document
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text.strip() for cell in row.cells))
            text = "\n".join(parts)
            if not text.strip():
                raise ValueError("The DOCX file contains no extractable text")
            return ExtractedDocument(text)
        except ImportError as exc:
            raise RuntimeError("DOCX support is not installed") from exc
    raise ValueError("Unsupported document type")


@lru_cache(maxsize=1)
def _get_ocr_reader():
    try:
        import easyocr

        return easyocr.Reader(["en"], gpu=False, verbose=False)
    except Exception as exc:
        raise RuntimeError("OCR could not be initialized") from exc


def _extract_pdf_pages_with_ocr(content: bytes, pages: list[str]) -> ExtractedDocument:
    scanned_pages = [
        page_number
        for page_number, page_text in enumerate(pages, 1)
        if len(re.sub(r"\s+", "", page_text)) < 30
    ]
    if not scanned_pages:
        return ExtractedDocument("\n".join(pages), pages=pages)

    try:
        reader = _get_ocr_reader()
        with tempfile.TemporaryDirectory() as directory:
            pdf_path = f"{directory}/document.pdf"
            with open(pdf_path, "wb") as pdf_file:
                pdf_file.write(content)
            confidence_values: list[float] = []
            for page_number in scanned_pages:
                image_prefix = f"{directory}/page-{page_number}"
                subprocess.run(
                    [
                        "pdftoppm",
                        "-png",
                        "-r",
                        "200",
                        "-f",
                        str(page_number),
                        "-l",
                        str(page_number),
                        "-singlefile",
                        pdf_path,
                        image_prefix,
                    ],
                    check=True,
                    capture_output=True,
                    text=True,
                )
                results = reader.readtext(f"{image_prefix}.png", detail=1)
                ocr_lines = [str(result[1]).strip() for result in results if str(result[1]).strip()]
                if not ocr_lines:
                    raise RuntimeError(f"OCR produced no text for PDF page {page_number}")
                pages[page_number - 1] = "\n".join(ocr_lines)
                confidence_values.extend(float(result[2]) for result in results if len(result) > 2)
    except FileNotFoundError as exc:
        raise RuntimeError("PDF rendering utility pdftoppm is not available") from exc
    except subprocess.CalledProcessError as exc:
        raise RuntimeError(f"PDF page rendering failed for page {scanned_pages[0]}") from exc
    except RuntimeError:
        raise
    except Exception as exc:
        raise RuntimeError("OCR processing failed") from exc

    text = "\n".join(pages)
    if not text.strip():
        raise RuntimeError("The PDF contains no extractable text after OCR")
    confidence = sum(confidence_values) / len(confidence_values) if confidence_values else None
    return ExtractedDocument(text, ocr_used=True, ocr_confidence=confidence, pages=pages)


def chunk_text(
    text: str,
    size: int = 1200,
    overlap: int = 150,
    pages: list[str] | None = None,
) -> list[dict[str, int | str | None]]:
    if pages is not None:
        chunks: list[dict[str, int | str | None]] = []
        chunk_index = 0
        for page_number, page_text in enumerate(pages, 1):
            page_chunks = chunk_text(page_text, size, overlap)
            for chunk in page_chunks:
                chunks.append({**chunk, "chunk_index": chunk_index, "page_number": page_number})
                chunk_index += 1
        return chunks
    normalized = re.sub(r"\s+", " ", text).strip()
    if not normalized:
        return []
    chunks: list[dict[str, int | str | None]] = []
    start = 0
    index = 0
    while start < len(normalized):
        end = min(len(normalized), start + size)
        chunks.append({"chunk_index": index, "content": normalized[start:end]})
        if end == len(normalized):
            break
        start = end - overlap
        index += 1
    return chunks
